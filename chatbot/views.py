from django.contrib import auth
from django.utils import timezone
from django.http import JsonResponse
from django.contrib.auth.models import User
from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.views.decorators.http import require_http_methods
from django.core.files.storage import default_storage
from django.conf import settings
import os
import json
import io

import google.generativeai as genai

from .models import Chat

# Configure Gemini API
GEMINI_API_KEY = "AIzaSyDBApMoTGxXL2GWPiV5_zjYYZB3MVQUijk"  
genai.configure(api_key=GEMINI_API_KEY)
model = genai.GenerativeModel("gemini-2.0-flash")

# ===================================
# PDF Text Extraction
# ===================================

def extract_text_from_pdf_pdfplumber(pdf_file):
    """Extract text from PDF using pdfplumber"""
    try:
        import pdfplumber
        pdf_file.seek(0)
        text = ""
        with pdfplumber.open(pdf_file) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
                if len(text) > 100000:
                    break
        return text.strip() if text.strip() else None
    except Exception as e:
        print(f"pdfplumber Error: {str(e)}")
        return None


def extract_text_from_pdf_pypdf2(pdf_file):
    """Extract text from PDF using PyPDF2"""
    try:
        import PyPDF2
        pdf_file.seek(0)
        pdf_reader = PyPDF2.PdfReader(pdf_file)
        text = ""
        for page_num in range(len(pdf_reader.pages)):
            page = pdf_reader.pages[page_num]
            extracted = page.extract_text()
            if extracted:
                text += extracted + "\n"
            if len(text) > 100000:
                break
        return text.strip() if text.strip() else None
    except Exception as e:
        print(f"PyPDF2 Error: {str(e)}")
        return None


def extract_pdf_content(pdf_file):
    """Extract text from PDF using multiple methods"""
    pdf_file.seek(0)
    
    # Try pdfplumber first
    text = extract_text_from_pdf_pdfplumber(pdf_file)
    if text and len(text) > 100:
        return text
    
    # Try PyPDF2
    pdf_file.seek(0)
    text = extract_text_from_pdf_pypdf2(pdf_file)
    if text and len(text) > 100:
        return text
    
    return None


# ===================================
# Document Text Extraction
# ===================================

def extract_text_from_docx(doc_file):
    """Extract text from DOCX files"""
    try:
        from docx import Document
        doc = Document(doc_file)
        text = ""
        
        for para in doc.paragraphs:
            if para.text.strip():
                text += para.text + "\n"
        
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text += cell.text + " | "
            text += "\n"
        
        return text.strip() if text.strip() else None
    except Exception as e:
        print(f"DOCX Error: {str(e)}")
        return None


def extract_text_from_txt(txt_file):
    """Extract text from TXT files"""
    try:
        txt_file.seek(0)
        content = txt_file.read().decode('utf-8', errors='ignore')
        return content.strip() if content.strip() else None
    except Exception as e:
        print(f"TXT Error: {str(e)}")
        return None


# ===================================
# Image Analysis with Gemini Vision
# ===================================

def analyze_image_with_gemini(image_file, user_query=""):
    """Analyze image using Gemini Vision API"""
    try:
        image_file.seek(0)
        image_data = image_file.read()
        
        filename = image_file.name.lower()
        if filename.endswith(('.jpg', '.jpeg')):
            mime_type = "image/jpeg"
        elif filename.endswith('.png'):
            mime_type = "image/png"
        elif filename.endswith('.gif'):
            mime_type = "image/gif"
        elif filename.endswith('.webp'):
            mime_type = "image/webp"
        else:
            mime_type = "image/jpeg"
        
        # Upload file to Gemini
        print(f"Uploading image: {filename}")
        uploaded_file = genai.upload_file(
            data=image_data,
            mime_type=mime_type
        )
        
        # Analyze image
        vision_model = genai.GenerativeModel("gemini-2.0-flash")
        
        if user_query:
            prompt = f"Here is an image. {user_query}\n\nAnalyze the image and answer the question."
        else:
            prompt = "Please analyze this image in detail and describe what you see."
        
        response = vision_model.generate_content([
            prompt,
            uploaded_file
        ])
        
        # Delete uploaded file
        genai.delete_file(uploaded_file.name)
        
        return response.text if response else "Could not analyze image"
    except Exception as e:
        print(f"Image Analysis Error: {str(e)}")
        return f"Error analyzing image: {str(e)}"


# ===================================
# File Content Extraction
# ===================================

def extract_file_content(file, user_query=""):
    """Extract content from uploaded file"""
    try:
        file_name = file.name.lower()
        content = None
        file_type = None
        
        if file_name.endswith('.pdf'):
            print(f"Processing PDF: {file.name}")
            content = extract_pdf_content(file)
            file_type = "PDF Document"
        
        elif file_name.endswith('.docx'):
            print(f"Processing DOCX: {file.name}")
            content = extract_text_from_docx(file)
            file_type = "Word Document (DOCX)"
        
        elif file_name.endswith('.doc'):
            print(f"Processing DOC: {file.name}")
            content = extract_text_from_docx(file)
            file_type = "Word Document (DOC)"
        
        elif file_name.endswith('.txt'):
            print(f"Processing TXT: {file.name}")
            content = extract_text_from_txt(file)
            file_type = "Text File"
        
        elif file_name.endswith(('.jpg', '.jpeg', '.png', '.gif', '.webp')):
            print(f"Processing Image: {file.name}")
            content = analyze_image_with_gemini(file, user_query)
            file_type = "Image"
        
        elif file_name.endswith(('.mp4', '.webm', '.mov', '.avi')):
            file_type = "Video File"
            content = f"Video file uploaded: {file_name}. Video content analysis is not yet supported."
        
        return content, file_type
    
    except Exception as e:
        print(f"File Processing Error: {str(e)}")
        return None, None


# ===================================
# AI Response Generation
# ===================================

def ask_gemini_with_files(message, file_contents=None):
    """Generate response from Gemini AI with file context"""
    try:
        if file_contents:
            # Build prompt with file contents
            prompt = f"""You are a helpful AI assistant answering questions about uploaded documents.

QUESTION: {message}

DOCUMENT CONTENT:
"""
            for file_info in file_contents:
                if file_info['content']:
                    prompt += f"\n[{file_info['file_type']}]\n"
                    prompt += file_info['content'][:8000]
                    prompt += "\n"
            
            prompt += f"\n\nProvide a clear, direct answer based on the document content above. Answer the question: {message}"
        else:
            # Regular chat without files
            prompt = f"User question: {message}\n\nProvide a helpful and direct answer."
        
        response = model.generate_content(prompt)
        
        if response and response.text:
            return response.text.strip()
        else:
            return "I couldn't generate a response. Please try again."
    
    except Exception as e:
        print(f"Gemini API Error: {str(e)}")
        return f"Error generating response: {str(e)}"


# ===================================
# Chatbot View
# ===================================

@login_required(login_url='login')
@require_http_methods(["GET", "POST"])
def chatbot(request):
    """Main chatbot view - GET: display history, POST: process message"""
    
    chats = Chat.objects.filter(user=request.user).order_by('created_at')
    
    if request.method == 'POST':
        message = request.POST.get('message', '').strip()
        files = request.FILES.getlist('files')
        
        if not message and not files:
            return JsonResponse({
                'success': False,
                'error': 'Please enter a message or upload a file'
            }, status=400)
        
        file_contents = []
        uploaded_info = []
        
        # Process files
        if files:
            for file in files:
                try:
                    # Validate file size
                    if file.size > 25 * 1024 * 1024:
                        return JsonResponse({
                            'success': False,
                            'error': f'File {file.name} exceeds 25MB limit'
                        }, status=400)
                    
                    print(f"\n[Processing] {file.name} ({file.size} bytes)")
                    
                    # Extract content
                    content, file_type = extract_file_content(file, message)
                    
                    if content:
                        file_contents.append({
                            'content': content,
                            'file_type': file_type
                        })
                        print(f"[Success] Extracted {len(content)} chars from {file.name}")
                    
                    uploaded_info.append({
                        'name': file.name,
                        'size': file.size,
                        'status': 'processed'
                    })
                
                except Exception as e:
                    print(f"[Error] {file.name}: {str(e)}")
                    return JsonResponse({
                        'success': False,
                        'error': f'Error processing {file.name}'
                    }, status=400)
        
        try:
            # Generate AI response
            response_text = ask_gemini_with_files(message, file_contents if file_contents else None)
            
            # Save to database
            chat = Chat(
                user=request.user,
                message=message,
                response=response_text,
                created_at=timezone.now()
            )
            chat.save()
            
            return JsonResponse({
                'success': True,
                'message': message,
                'response': response_text,
                'files_processed': len(uploaded_info)
            })
        
        except Exception as e:
            print(f"[Error] Response generation: {str(e)}")
            return JsonResponse({
                'success': False,
                'error': 'Failed to generate response'
            }, status=500)
    
    return render(request, 'chatbot.html', {
        'chats': chats,
        'total_chats': chats.count()
    })


# ===================================
# Authentication Views
# ===================================

@require_http_methods(["GET", "POST"])
def login(request):
    """User login"""
    if request.user.is_authenticated:
        return redirect('chatbot')
    
    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        password = request.POST.get('password', '').strip()
        
        if not username or not password:
            return render(request, 'login.html', {
                'error_message': 'Username and password required'
            })
        
        user = auth.authenticate(request, username=username, password=password)
        
        if user:
            auth.login(request, user)
            return redirect('chatbot')
        else:
            return render(request, 'login.html', {
                'error_message': 'Invalid credentials'
            })
    
    return render(request, 'login.html')


@require_http_methods(["GET", "POST"])
def register(request):
    """User registration"""
    if request.user.is_authenticated:
        return redirect('chatbot')
    
    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        email = request.POST.get('email', '').strip()
        password1 = request.POST.get('password1', '').strip()
        password2 = request.POST.get('password2', '').strip()
        
        if not all([username, email, password1, password2]):
            return render(request, 'register.html', {
                'error_message': 'All fields required'
            })
        
        if password1 != password2:
            return render(request, 'register.html', {
                'error_message': 'Passwords do not match'
            })
        
        if len(password1) < 6:
            return render(request, 'register.html', {
                'error_message': 'Password must be 6+ characters'
            })
        
        if User.objects.filter(username=username).exists():
            return render(request, 'register.html', {
                'error_message': 'Username already taken'
            })
        
        if User.objects.filter(email=email).exists():
            return render(request, 'register.html', {
                'error_message': 'Email already registered'
            })
        
        try:
            user = User.objects.create_user(
                username=username,
                email=email,
                password=password1
            )
            auth.login(request, user)
            return redirect('chatbot')
        except Exception as e:
            return render(request, 'register.html', {
                'error_message': 'Account creation failed'
            })
    
    return render(request, 'register.html')


@login_required(login_url='login')
def logout(request):
    """User logout"""
    auth.logout(request)
    return redirect('login')


# ===================================
# API Endpoints
# ===================================

@login_required(login_url='login')
def get_chat_history(request):
    """Get chat history as JSON"""
    chats = Chat.objects.filter(user=request.user).order_by('-created_at')[:50]
    
    return JsonResponse({
        'chats': [
            {
                'id': chat.id,
                'message': chat.message,
                'response': chat.response,
                'created_at': chat.created_at.isoformat()
            }
            for chat in chats
        ]
    })


@login_required(login_url='login')
@require_http_methods(["POST"])
def clear_chat_history(request):
    """Clear all chats"""
    try:
        Chat.objects.filter(user=request.user).delete()
        return JsonResponse({
            'success': True,
            'message': 'Chat history cleared'
        })
    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=400)


@login_required(login_url='login')
@require_http_methods(["POST"])
def delete_chat(request, chat_id):
    """Delete specific chat"""
    try:
        chat = Chat.objects.get(id=chat_id, user=request.user)
        chat.delete()
        return JsonResponse({
            'success': True,
            'message': 'Chat deleted'
        })
    except Chat.DoesNotExist:
        return JsonResponse({
            'success': False,
            'error': 'Not found'
        }, status=404)
    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=400)


@login_required(login_url='login')
def search_chats(request):
    """Search chat history"""
    query = request.GET.get('q', '').strip()
    
    if not query:
        return JsonResponse({'chats': []})
    
    chats = Chat.objects.filter(
        user=request.user,
        message__icontains=query
    ).order_by('-created_at')[:20]
    
    return JsonResponse({
        'chats': [
            {
                'id': chat.id,
                'message': chat.message[:100],
                'created_at': chat.created_at.isoformat()
            }
            for chat in chats
        ]
    })