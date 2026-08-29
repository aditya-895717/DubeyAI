"""Document text extraction for chat file uploads.

Every extractor returns plain text and raises DocumentExtractionError with a
user-safe message on failure — callers surface that text directly, so it must
never contain a traceback or a filesystem path.

OCR note: image extraction depends on `pytesseract` AND the `tesseract` system
binary. Neither is available on Vercel's read-only Python serverless runtime, so
the image path degrades to a clear "OCR unavailable" message rather than a 500.
"""

import logging
import os

logger = logging.getLogger(__name__)


class DocumentExtractionError(Exception):
    """A safe, user-facing document extraction failure."""


# Keep in sync with the `accept` attribute on the chat file input.
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".png", ".jpg", ".jpeg"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}

MAX_UPLOAD_BYTES = 15 * 1024 * 1024  # 15 MB

# Roughly 8,000 tokens at ~4 characters per token. Extracted text is truncated
# to this before it is ever sent to a provider, so a 300-page PDF cannot blow
# the model's context window (or the bill).
MAX_EXTRACTED_CHARS = 32_000

TRUNCATION_NOTICE = (
    "\n\n[... document truncated — only the first portion is available "
    "as context ...]"
)


def get_extension(filename):
    return os.path.splitext(filename or "")[1].lower()


def validate_upload(uploaded_file):
    """Raise DocumentExtractionError unless `uploaded_file` is acceptable."""
    extension = get_extension(uploaded_file.name)
    if extension not in ALLOWED_EXTENSIONS:
        allowed = ", ".join(sorted(ALLOWED_EXTENSIONS))
        raise DocumentExtractionError(
            f"Unsupported file type '{extension or 'unknown'}'. Allowed: {allowed}."
        )
    if uploaded_file.size == 0:
        raise DocumentExtractionError("That file is empty.")
    if uploaded_file.size > MAX_UPLOAD_BYTES:
        limit_mb = MAX_UPLOAD_BYTES // (1024 * 1024)
        raise DocumentExtractionError(f"File is too large. The limit is {limit_mb} MB.")
    return extension


def truncate(text):
    """Cap extracted text so it always fits inside the AI context window."""
    text = (text or "").strip()
    if len(text) <= MAX_EXTRACTED_CHARS:
        return text
    return text[:MAX_EXTRACTED_CHARS].rstrip() + TRUNCATION_NOTICE


# ---------------------------------------------------------------------------
# Per-format extractors
# ---------------------------------------------------------------------------

def _extract_pdf(file_obj):
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover - dependency is in requirements
        raise DocumentExtractionError("PDF support is not installed on the server.") from exc

    pages = []
    try:
        with pdfplumber.open(file_obj) as pdf:
            for page in pdf.pages:
                pages.append(page.extract_text() or "")
                # Stop early once we already have more text than we can send.
                if sum(len(part) for part in pages) > MAX_EXTRACTED_CHARS:
                    break
    except DocumentExtractionError:
        raise
    except Exception as exc:
        logger.warning("pdfplumber failed, trying PyPDF2 fallback: %s", exc)
        return _extract_pdf_fallback(file_obj)

    text = "\n\n".join(part for part in pages if part.strip())
    if not text.strip():
        # A scanned PDF has no text layer at all — pdfplumber returns empty
        # strings rather than raising, so this needs an explicit check.
        raise DocumentExtractionError(
            "This PDF has no selectable text — it looks like a scan or images "
            "only. Try a text-based PDF, or export it with OCR first."
        )
    return text


def _extract_pdf_fallback(file_obj):
    try:
        from PyPDF2 import PdfReader
    except ImportError as exc:
        raise DocumentExtractionError("Could not read that PDF. It may be corrupted or encrypted.") from exc

    try:
        file_obj.seek(0)
        reader = PdfReader(file_obj)
        text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception as exc:
        raise DocumentExtractionError("Could not read that PDF. It may be corrupted or encrypted.") from exc

    if not text.strip():
        raise DocumentExtractionError(
            "This PDF has no selectable text — it looks like a scan or images only."
        )
    return text


def _extract_docx(file_obj):
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - dependency is in requirements
        raise DocumentExtractionError("DOCX support is not installed on the server.") from exc

    try:
        document = docx.Document(file_obj)
    except Exception as exc:
        raise DocumentExtractionError("Could not read that DOCX file. It may be corrupted.") from exc

    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    # Tables carry most of the useful content in report-style documents, so
    # flatten them into pipe-separated rows rather than dropping them.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    text = "\n".join(parts)
    if not text.strip():
        raise DocumentExtractionError("That DOCX file contains no readable text.")
    return text


def _extract_txt(file_obj):
    raw = file_obj.read()
    if isinstance(raw, str):
        text = raw
    else:
        # UTF-16 is only tried on an explicit BOM: without one it happily
        # decodes arbitrary byte pairs into mojibake and would shadow latin-1.
        if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
            encodings = ("utf-16", "utf-8", "latin-1")
        else:
            encodings = ("utf-8-sig", "utf-8", "latin-1")

        for encoding in encodings:
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise DocumentExtractionError("Could not decode that text file.")
    if not text.strip():
        raise DocumentExtractionError("That text file is empty.")
    return text


def _extract_image(file_obj):
    """OCR an image. Requires pytesseract *and* the tesseract system binary."""
    try:
        import pytesseract
        from PIL import Image
    except ImportError as exc:
        raise DocumentExtractionError(
            "Image OCR is not available on this server. Please upload a PDF, "
            "DOCX, or TXT file instead."
        ) from exc

    try:
        text = pytesseract.image_to_string(Image.open(file_obj))
    except Exception as exc:
        # Most commonly TesseractNotFoundError — the Python package installs
        # fine but the native binary is missing (e.g. on serverless hosts).
        logger.warning("OCR failed: %s", exc)
        raise DocumentExtractionError(
            "Image OCR is not available on this server. Please upload a PDF, "
            "DOCX, or TXT file instead."
        ) from exc

    if not text.strip():
        raise DocumentExtractionError("No readable text was found in that image.")
    return text


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".txt": _extract_txt,
    ".png": _extract_image,
    ".jpg": _extract_image,
    ".jpeg": _extract_image,
}


def extract_text(file_obj, extension):
    """Extract and truncate text from an already-validated upload."""
    extractor = _EXTRACTORS.get(extension)
    if extractor is None:
        raise DocumentExtractionError(f"Unsupported file type '{extension}'.")

    try:
        file_obj.seek(0)
    except (AttributeError, OSError):
        pass

    text = extractor(file_obj)
    return truncate(text)
